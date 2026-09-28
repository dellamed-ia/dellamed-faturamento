"""Gera build/data.json a partir da planilha Dados_dellamed_mix_e_clientes.xlsx.

Serve para publicar a primeira versão antes da conexão com o Qlik, ou como
plano B se o Qlik estiver fora do ar.

Uso:
  python scripts/from_xlsx.py Dados_dellamed_mix_e_clientes.xlsx --periodo "Jan–Ago 2026"
"""
import argparse, json
from datetime import datetime, timezone
from pathlib import Path
import openpyxl


def num(v):
    return v if isinstance(v, (int, float)) else 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--periodo", required=True)
    ap.add_argument("--out", default="build/data.json")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx, data_only=True, read_only=True)
    rows = []
    for r in wb["mix produtos"].iter_rows(min_row=2, values_only=True):
        if r[0] and r[1] and r[2]:
            rows.append([r[0], r[1], r[2], num(r[3]), round(num(r[5]), 2)])

    cli, soma = {}, 0
    for r in wb["segmentos"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            cli[f"|{r[0]}|"] = int(num(r[1]))
            soma += int(num(r[1]))
    cli["||"] = soma

    dados = {
        "meta": {
            "periodo": a.periodo,
            "atualizado": datetime.now(timezone.utc).isoformat(),
            "fonte": "Planilha " + Path(a.xlsx).name,
            "cliAprox": True,  # total = soma dos segmentos (pode contar o mesmo cliente 2x)
        },
        "rows": rows,
        "cli": cli,
    }
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text(json.dumps(dados, ensure_ascii=False), encoding="utf-8")
    print(f"OK: {len(rows)} registros -> {a.out}")


if __name__ == "__main__":
    main()
